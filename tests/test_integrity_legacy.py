"""Exact, expiring legacy acceptance; real gate and export paths, no mocks."""
import copy
import csv
import gzip
import hashlib
import importlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from src import integrity_audit as A, integrity_triage as T
from src import master as M

REPO = Path(__file__).resolve().parent.parent
REVIEWED = '2026-10-26T00:00:00Z'
EXPIRES = '2026-11-25T00:00:00Z'
NOW = datetime(2026,10,26,12,tzinfo=timezone.utc)
SCOPE_PATH = 'docs/evidence/source-integrity-legacy-scope-20261026.json.gz'
POLICY_PATH = 'docs/evidence/source-integrity-legacy-retention-policy-proposal-20261026.json'


def module(case):
    case.assertIsNotNone(importlib.util.find_spec('src.integrity_legacy'),'legacy policy implementation missing')
    return importlib.import_module('src.integrity_legacy')


def sample():
    pack=json.loads(gzip.decompress((REPO/SCOPE_PATH).read_bytes()))
    return copy.deepcopy(next(i for i in pack['items'] if i['classification']=='ORIGIN_UNRESOLVED'
        and i['finding']['master']['sources']=='OFAC' and not i['finding']['master']['categories']
        and not i['index_links'] and not i['evidence']))


def scope_fixture():
    i=sample();return {i['review_id']:i}


def active(draft):
    ledger=copy.deepcopy(draft)
    ledger.update(state='ACTIVE',reviewed_by='fixture reviewer',reason='Explicit unresolved screening-string retention')
    return ledger


def write_master(root,rows):
    path=root/'data/master/master.csv';path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=M.FIELDS);writer.writeheader();writer.writerows(rows)


class ScopeTests(unittest.TestCase):
    def test_real_scope_is_fixed_and_old_review_ids_remain_stable(self):
        L=module(self);scope=L.load_scope(REPO)
        self.assertEqual(len(scope),2429)
        self.assertEqual({i['classification'] for i in scope.values()},
            {'ORIGIN_UNRESOLVED','NOTICE_ALIAS_FIELD_REVIEW','NOTICE_MENTION_ONLY','NON_NAME_FIELD_ONLY','WEAK_EVIDENCE_CONFLICT'})
        for i in scope.values():
            self.assertEqual(T.review_id(i['finding'],i['classification'],i['index_links'],i['evidence']),i['review_id'])

    def test_tampered_or_missing_scope_body_and_unsafe_link_refuse_authority(self):
        L=module(self)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            for rel in [SCOPE_PATH,POLICY_PATH]:
                p=root/rel;p.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(REPO/rel,p)
            self.assertEqual(len(L.load_scope(root)),2429)
            p=root/SCOPE_PATH;p.write_bytes(gzip.compress(b'{}'))
            with self.assertRaises(ValueError):L.load_scope(root)
            p.unlink();p.symlink_to(REPO/SCOPE_PATH)
            with self.assertRaises(ValueError):L.load_scope(root)


class PreparationTests(unittest.TestCase):
    def setUp(self):
        self.L=module(self);self.item=sample();self.row=self.item['finding']['master'];self.scope=scope_fixture()

    def test_preparation_preserves_names_status_high_risk_and_existing_fields_without_mutation(self):
        before=copy.deepcopy(self.row)
        rows,ledger=self.L.prepare([self.row],self.scope,REVIEWED,EXPIRES)
        self.assertEqual(self.row,before);self.assertEqual(len(rows),1)
        for field in ['match_key','display_name','status','risk_level','sources','categories','first_seen_ms','last_updated_ms','variants']:
            self.assertEqual(rows[0][field],before[field],field)
        self.assertTrue(rows[0]['risk_type'].startswith('旧登録・一次出典未確認'))
        self.assertIn(before['risk_type'],rows[0]['risk_type'])
        self.assertIn('LEGACY_UNVERIFIED_V1:ORIGIN_UNRESOLVED',rows[0]['review_flag'])
        self.assertIn(before['remark'],rows[0]['remark']);self.assertIn(EXPIRES,rows[0]['remark'])
        self.assertEqual(ledger['state'],'DRAFT');self.assertIsNone(ledger['reviewed_by'])
        self.assertNotEqual(ledger['entries'][0]['accepted_review_id'],self.item['review_id'])

    def test_existing_review_flags_and_remarks_are_appended_not_replaced(self):
        row=dict(self.row,review_flag='KEEP_PRIOR_FLAG',remark='prior remark')
        annotated=self.L.annotate(row,'ORIGIN_UNRESOLVED',EXPIRES)
        self.assertTrue(annotated['review_flag'].startswith('KEEP_PRIOR_FLAG;'))
        self.assertTrue(annotated['remark'].startswith('prior remark\n'))
        self.assertEqual(row['review_flag'],'KEEP_PRIOR_FLAG')

    def test_changed_original_duplicate_missing_or_extra_scope_id_cannot_be_prepared(self):
        cases=[[dict(self.row,risk_level='低')],[self.row,self.row],[]]
        for rows in cases:
            with self.subTest(rows=len(rows)):
                with self.assertRaises(ValueError):self.L.prepare(rows,self.scope,REVIEWED,EXPIRES)
        with self.assertRaises(ValueError):self.L.prepare([self.row],self.scope,REVIEWED,EXPIRES,review_ids=['0'*64])

    def test_term_is_at_most_30_days_with_strict_utc_dates(self):
        for expires in ['2026-11-25T00:00:01Z','2026-10-26T00:00:00Z','2026-11-25','2026-11-25T00:00:00+00:00']:
            with self.subTest(expires=expires):
                with self.assertRaises(ValueError):self.L.prepare([self.row],self.scope,REVIEWED,expires)

    def test_real_preparation_cli_creates_only_drafts_outside_business_data_and_refuses_overwrite(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            for rel in [SCOPE_PATH,POLICY_PATH]:
                p=root/rel;p.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(REPO/rel,p)
            write_master(root,[self.row]);before=(root/'data/master/master.csv').read_bytes()
            cmd=[sys.executable,'-m','src.integrity_legacy','--root',str(root),'--prepare','--reviewed-at',REVIEWED,
                 '--expires-at',EXPIRES,'--review-id',self.item['review_id'],'--out-dir',str(root/'candidate')]
            p=subprocess.run(cmd,cwd=REPO,capture_output=True,text=True)
            self.assertEqual(p.returncode,0,p.stderr)
            ledger=json.loads((root/'candidate/draft-legacy-acceptances.json').read_text())
            self.assertEqual(ledger['state'],'DRAFT');self.assertFalse((root/'data/review/integrity_legacy_acceptances.json').exists())
            self.assertEqual(before,(root/'data/master/master.csv').read_bytes())
            self.assertNotEqual(subprocess.run(cmd,cwd=REPO,capture_output=True).returncode,0)
            bad=cmd[:-1]+[str(root/'data/candidate')]
            self.assertNotEqual(subprocess.run(bad,cwd=REPO,capture_output=True).returncode,0)


class GateTests(unittest.TestCase):
    def setUp(self):
        self.L=module(self)
        from tests.test_integrity_publication import GateTests as G
        self.gate=G();self.gate.setUp();self.addCleanup(self.gate.doCleanups);self.root=self.gate.root
        self.item=sample();self.scope=scope_fixture()
        with self.gate.mpath.open() as stream:rows=list(csv.DictReader(stream))
        rows=[{f:r.get(f,'') for f in M.FIELDS} for r in rows]+[copy.deepcopy(self.item['finding']['master'])]
        self.original=copy.deepcopy(rows)
        self.rows,self.draft=self.L.prepare(rows,self.scope,REVIEWED,EXPIRES)
        self.ledger=active(self.draft)
        for rel in [SCOPE_PATH,POLICY_PATH]:
            p=self.root/rel;p.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(REPO/rel,p)
        self.install()

    def install(self,ledger=None,rows=None):
        write_master(self.root,rows if rows is not None else self.rows)
        p=self.root/'data/review/integrity_legacy_acceptances.json';p.parent.mkdir(exist_ok=True)
        p.write_text(json.dumps(ledger if ledger is not None else self.ledger))
        from src import dashboard as D
        D.write_list(self.root,rows if rows is not None else self.rows)
        D.write_screening_gzip(self.root,rows if rows is not None else self.rows)

    def report(self):
        paths=[p for p in self.root.glob('data/integrity_audit/*.json.gz') if not p.name.startswith('triage-')]
        return json.loads(gzip.decompress(sorted(paths)[-1].read_bytes()))

    def test_exact_scope_annotations_and_projections_allow_review_without_mutating_master(self):
        before=self.gate.mpath.read_bytes();p=self.gate.run_gate('success','mof')
        self.assertEqual(p.returncode,0,p.stdout+p.stderr);self.assertTrue((self.root/'published').exists())
        self.assertEqual(before,self.gate.mpath.read_bytes())
        r=self.report();self.assertEqual(r['status'],'REVIEW')
        self.assertEqual(sum(f.get('legacy_accepted',False) for f in r['findings']),1)
        triage=json.loads(gzip.decompress(next(self.root.glob('data/integrity_audit/triage-*.json.gz')).read_bytes()))
        self.assertEqual(triage['legacy_accepted_count'],1);self.assertEqual(triage['approved_count'],0)

    def test_draft_ledger_and_unmarked_or_changed_master_never_release(self):
        for ledger,rows in [(self.draft,self.rows),(self.ledger,self.original),
                            (self.ledger,[dict(r,risk_level='低') if r['match_key']==self.item['master_key'] else r for r in self.rows])]:
            with self.subTest(state=ledger['state'],marked=rows==self.rows):
                self.install(ledger,rows);p=self.gate.run_gate('success','ofac')
                self.assertNotEqual(p.returncode,0);self.assertFalse((self.root/'published').exists())
                self.assertTrue(any(f['kind'].startswith('LEGACY_') for f in self.report()['findings']))

    def test_expiration_boundary_and_future_review_do_not_release(self):
        r=A.audit(self.root,['ofac','mof']);t=T.triage(self.root,r)
        # Rebuild an unaccepted report for the pure finalizer using actual finding content.
        for f in r['findings']:
            if f.get('legacy_accepted'):f.pop('legacy_accepted');f['blocking']=True
        t=T.triage(self.root,r)
        for now in [datetime(2026,11,25,tzinfo=timezone.utc),datetime(2026,10,25,tzinfo=timezone.utc)]:
            with self.subTest(now=now):
                result=self.L.apply(self.root,r,t,self.ledger,self.scope,now=now)
                self.assertTrue(any(f.get('blocking',True) for f in result['findings']))

    def test_valid_legacy_approval_cannot_waive_missing_alias_ids_failed_fetch_or_raw_hash(self):
        # Every variant exercises the actual bash gate and retains failure evidence.
        baseline=self.gate.mpath.read_bytes()
        with self.gate.mpath.open() as s:rows=list(csv.DictReader(s))
        write_master(self.root,[r for r in rows if r['display_name']!='Jane Other'])
        p=self.gate.run_gate('success','mof');self.assertNotEqual(p.returncode,0);self.assertIn('MISSING_ALIAS',p.stdout)
        self.gate.mpath.write_bytes(baseline)
        p=self.gate.run_gate('failure','ofac');self.assertNotEqual(p.returncode,0);self.assertIn('UPSTREAM_FAILED',p.stdout)
        raw=self.root/'data/raw/ofac_sdn/prim.gz';raw.write_bytes(gzip.compress(b'changed'))
        p=self.gate.run_gate('success','ofac');self.assertNotEqual(p.returncode,0);self.assertIn('UNDETERMINED',p.stdout)
        self.assertFalse((self.root/'published').exists())

    def test_current_party_id_mismatch_remains_a_veto(self):
        p=self.root/'data/master/ofac_alias_history.csv'
        with p.open() as stream:reader=csv.DictReader(stream);fields=reader.fieldnames;rows=list(reader)
        rows[0]['party_id']='999999'
        with p.open('w',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=fields);writer.writeheader();writer.writerows(rows)
        result=self.gate.run_gate('success','mof')
        self.assertNotEqual(result.returncode,0);self.assertIn('SOURCE_ID_MISMATCH',result.stdout)
        self.assertFalse((self.root/'published').exists())

    def test_duplicate_unmarked_search_copy_cannot_coexist_with_annotated_row(self):
        row=self.item['finding']['master'];path=self.root/'data/dashboard/list.csv'
        with path.open('a',newline='') as stream:
            csv.writer(stream).writerow([row['display_name'],row['risk_type'],row['status'],row['risk_level']])
        result=self.gate.run_gate('success','ofac')
        self.assertNotEqual(result.returncode,0)
        self.assertIn('LEGACY_DISPLAY_MISMATCH',self.report()['counts'])
        self.assertFalse((self.root/'published').exists())

    def test_extra_unmarked_screening_copy_cannot_hide_under_another_key(self):
        path=self.root/'data/dashboard/screening.csv.gz'
        reader=csv.DictReader(__import__('io').StringIO(gzip.decompress(path.read_bytes()).decode()))
        fields=reader.fieldnames;rows=list(reader)
        row=copy.deepcopy(next(r for r in rows if r['match_key']==self.item['master_key']))
        row.update(match_key='stale-copy',リスクタイプ=self.item['finding']['master']['risk_type'],要確認='')
        stream=__import__('io').StringIO(newline='');writer=csv.DictWriter(stream,fieldnames=fields)
        writer.writeheader();writer.writerows([row]+rows);path.write_bytes(gzip.compress(stream.getvalue().encode()))
        result=self.gate.run_gate('success','mof')
        self.assertNotEqual(result.returncode,0)
        self.assertIn('LEGACY_DISPLAY_MISMATCH',self.report()['counts'])
        self.assertFalse((self.root/'published').exists())

    def test_corrupt_screening_gzip_keeps_failure_evidence(self):
        (self.root/'data/dashboard/screening.csv.gz').write_bytes(b'not gzip')
        result=self.gate.run_gate('success','ofac')
        self.assertNotEqual(result.returncode,0)
        self.assertIn('LEGACY_DISPLAY_MISMATCH',self.report()['counts'])
        self.assertFalse((self.root/'published').exists())

    def test_expired_acceptance_and_missing_ledger_refuse_export_validation(self):
        for now in [datetime(2026,11,25,tzinfo=timezone.utc),datetime(2026,10,25,tzinfo=timezone.utc)]:
            with self.assertRaises(ValueError):self.L.validate_for_export(self.root,self.rows,now=now)
        (self.root/'data/review/integrity_legacy_acceptances.json').unlink()
        with self.assertRaises(ValueError):self.L.validate_for_export(self.root,self.rows,now=NOW)

    def test_ordinary_retention_does_not_approve_legacy(self):
        from src import integrity_decisions as D
        report=A.audit(self.root,['ofac','mof']);triage=T.triage(self.root,report)
        item=next(i for i in triage['items'] if i['classification']=='ORIGIN_UNRESOLVED')
        # A separate ordinary approval cannot manufacture eligibility.
        for finding in report['findings']:
            finding.pop('legacy_accepted',None);finding['blocking']=True
        triage=T.triage(self.root,report)
        ledger=dict(schema_version=1,entries=[dict(review_id=item['review_id'],decision='APPROVED_RETAIN',
            reviewed_by='fixture',reason='fixture',reviewed_at=REVIEWED,expires_at=EXPIRES,
            reviewed_evidence_sha256=item['evidence_sha256'])])
        result=D.apply(report,triage,ledger,now=NOW)
        self.assertFalse(any(f.get('retention_approved') for f in result['findings']))
        self.assertEqual(result['retention_review'][0]['status'],'INELIGIBLE')

    def test_stale_search_display_or_missing_screening_warning_blocks_publication(self):
        p=self.root/'data/dashboard/list.csv';original=p.read_bytes()
        p.write_text(p.read_text().replace('旧登録・一次出典未確認',''))
        result=self.gate.run_gate('success','ofac');self.assertNotEqual(result.returncode,0)
        self.assertFalse((self.root/'published').exists());p.write_bytes(original)
        path=self.root/'data/dashboard/screening.csv.gz';path.unlink()
        result=self.gate.run_gate('success','mof');self.assertNotEqual(result.returncode,0)
        self.assertFalse((self.root/'published').exists())

    def test_new_relation_no_longer_matches_exact_frozen_review_id(self):
        directory=self.root/'docs/evidence/integrity-notices';directory.mkdir(parents=True)
        body=('<div class="field--name-field-body"><p>'+self.item['display_name']+'</p></div>').encode()
        raw=directory/'notice.html.gz';raw.write_bytes(gzip.compress(body,mtime=0))
        (directory/'manifest.json').write_text(json.dumps([dict(raw_path=raw.relative_to(self.root).as_posix(),
            source_url='https://ofac.treasury.gov/recent-actions/20250630',source_hash=hashlib.sha256(body).hexdigest())]))
        p=self.gate.run_gate('success','ofac');self.assertNotEqual(p.returncode,0)
        self.assertFalse((self.root/'published').exists())

    def test_outside_scope_unknown_decision_duplicate_and_overlong_term_fail_closed(self):
        variants=[]
        x=copy.deepcopy(self.ledger);x['entries'][0]['original_review_id']='0'*64;variants.append(x)
        x=copy.deepcopy(self.ledger);x['decision']='APPROVED_RETAIN';variants.append(x)
        x=copy.deepcopy(self.ledger);x['entries']*=2;variants.append(x)
        x=copy.deepcopy(self.ledger);x['expires_at']='2026-11-25T00:00:01Z';variants.append(x)
        for ledger in variants:
            with self.subTest(ledger=ledger.get('decision')):
                self.install(ledger);p=self.gate.run_gate('success','mof')
                self.assertNotEqual(p.returncode,0);self.assertFalse((self.root/'published').exists())

    def test_diff_excel_warning_and_draft_veto_even_without_diff(self):
        from openpyxl import load_workbook
        diff=self.root/'diff.csv';out=self.root/'diff.xlsx'
        with diff.open('w',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=['種別','受取人名','変更前','変更後']);writer.writeheader()
            writer.writerow({'種別':'変更','受取人名':self.item['display_name'],'変更前':'a','変更後':'b'})
        cmd=[sys.executable,'-m','src.export_diff_excel','--master',str(self.gate.mpath),
            '--diff',str(diff),'--out',str(out)]
        result=subprocess.run(cmd,cwd=REPO,capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        ws=load_workbook(out)['Sheet1'];self.assertEqual(ws.max_column,7)
        self.assertIn('旧登録・一次出典未確認',ws.cell(2,2).value);self.assertIn(EXPIRES,ws.cell(2,7).value)
        before=out.read_bytes();self.install(self.draft);diff.unlink()
        result=subprocess.run(cmd,cwd=REPO,capture_output=True,text=True)
        self.assertNotEqual(result.returncode,0);self.assertEqual(out.read_bytes(),before)

    def test_excel_cli_preserves_seven_columns_and_warns_without_apps_script_change(self):
        from openpyxl import load_workbook
        out=self.root/'book.xlsx'
        p=subprocess.run([sys.executable,'-m','src.export_excel','--indir',str(self.root/'data/master'),
                          '--out',str(out)],cwd=REPO,capture_output=True,text=True)
        self.assertEqual(p.returncode,0,p.stderr)
        wb=load_workbook(out);ws=wb['Sheet1'];self.assertEqual(ws.max_column,7)
        row=next(values for values in ws.iter_rows(min_row=2,values_only=True) if values[0]==self.item['display_name'])
        self.assertIn('旧登録・一次出典未確認',row[1]);self.assertIn(EXPIRES,row[6])
        self.assertEqual(row[2],'有効');self.assertEqual(row[3],self.item['finding']['master']['risk_level'])
        before=out.read_bytes();self.install(self.draft)
        p=subprocess.run([sys.executable,'-m','src.export_excel','--indir',str(self.root/'data/master'),
                          '--out',str(self.root/'rejected.xlsx')],cwd=REPO,capture_output=True,text=True)
        self.assertNotEqual(p.returncode,0);self.assertFalse((self.root/'rejected.xlsx').exists())
        self.assertEqual(out.read_bytes(),before)


if __name__=='__main__':unittest.main()
